"""PRIORITY — who to call first, and why, with no model call.

A profile's ``priority.yaml`` lists plain rules over what Harvey already
knows about each company: the signals its site gave (observations), the
notes imported with it (CRM status, what was said on a call), its industry,
whether it has a site at all. Each rule that matches adds or takes points and
leaves its reason, so a ranking is always explainable line by line:

    base: 50
    rules:
      - if: {signal: RUNNING_GOOGLE_ADS, value: 1}
        points: 10
        reason: "Already pays for ads: there is an acquisition budget"
      - if: {notes: "no (le )?interesa"}
        points: -40
        reason: "Said on the call it isn't interested"

Conditions (all must hold): ``signal`` (+ ``value``/``min``/``max`` on its
number, ``text`` regex on its text), ``absent`` (profiled, signal not seen),
``notes`` and ``industry`` (regexes), ``no_website`` and ``unreadable``
(booleans). A rule with ``product`` names the offer for the companies it
matches (first such rule wins). ``{text}`` and ``{value}`` in a reason are
filled from the matched signal.
"""

import json
import re
from dataclasses import dataclass, field
from pathlib import Path

import aiosqlite
import yaml

from harvey.paths import PROFILE_DIR, PROJECT_ROOT

RULES_FILE = "priority.yaml"


def rules_path() -> Path:
    for base in (PROFILE_DIR, PROJECT_ROOT):
        path = base / RULES_FILE
        if path.is_file():
            return path
    raise FileNotFoundError(
        f"No {RULES_FILE} in {PROFILE_DIR}: write the rules that decide who goes first "
        "(see harvey/priority.py for the format).")


def load_rules(path: Path | None = None) -> dict:
    data = yaml.safe_load((path or rules_path()).read_text(encoding="utf-8")) or {}
    data.setdefault("base", 50)
    data.setdefault("rules", [])
    data.setdefault("default_product", "")
    return data


@dataclass
class Ranked:
    company: dict
    score: int
    product: str
    reasons: list[str] = field(default_factory=list)
    ask_for: str = ""
    decision_maker: bool = False      # a named decision maker is already known


def _num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _matches(cond: dict, company: dict, obs: dict[str, list[dict]]) -> tuple[bool, dict]:
    """(matched, the observation that matched, if any)."""
    hit: dict = {}
    profiled = company.get("_profiled", False)
    if "no_website" in cond and bool(cond["no_website"]) != (not company.get("domain")):
        return False, hit
    if "unreadable" in cond:
        unreadable = bool(company.get("domain")) and profiled and not obs
        if bool(cond["unreadable"]) != unreadable:
            return False, hit
    if "notes" in cond and not re.search(cond["notes"], company.get("notes") or "", re.I):
        return False, hit
    if "industry" in cond and not re.search(cond["industry"], company.get("industry") or "", re.I):
        return False, hit
    if "absent" in cond:
        if not profiled or not obs or cond["absent"] in obs:
            return False, hit
    if "signal" in cond:
        candidates = obs.get(cond["signal"], [])
        for o in candidates:
            value = _num(o.get("value_num"))
            if "value" in cond and value != _num(cond["value"]):
                continue
            if "min" in cond and (value is None or value < _num(cond["min"])):
                continue
            if "max" in cond and (value is None or value > _num(cond["max"])):
                continue
            if "text" in cond and not re.search(cond["text"], o.get("value_text") or "", re.I):
                continue
            hit = o
            break
        else:
            return False, hit
    return True, hit


def _fmt(reason: str, hit: dict) -> str:
    value = hit.get("value_num")
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = hit.get("value_text") or ""
    return reason.replace("{text}", str(text)).replace("{value}", "" if value is None else str(value))


def _ask_for(company: dict, obs: dict[str, list[dict]]) -> str:
    from harvey.collectors.profile import decision_weight

    def weight(o):
        title = (o.get("value_text") or "").split(" — ", 1)[-1]
        return decision_weight(title)

    for code in ("DECISION_MAKER_TITLE", "CONTACT_FOUND"):
        if obs.get(code):
            # The owner before a department head; the site's own page before a search hit.
            best = max(obs[code], key=lambda o: (weight(o), "(según" not in (o.get("value_text") or "")))
            return best.get("value_text") or ""
    m = re.search(r"^Contacto: (.+)$", company.get("notes") or "", re.M)
    return f"{m.group(1)} (del CRM)" if m else ""


def score_company(company: dict, obs: dict[str, list[dict]], rules: dict) -> Ranked:
    score = float(rules["base"])
    product = ""
    reasons = []
    for rule in rules["rules"]:
        ok, hit = _matches(rule.get("if") or {}, company, obs)
        if not ok:
            continue
        points = float(rule.get("points") or 0)
        score += points
        if rule.get("product") and not product:
            product = rule["product"]
        if rule.get("reason"):
            sign = f"{points:+.0f} " if points else ""
            reasons.append(f"{sign}{_fmt(rule['reason'], hit)}".strip())
    return Ranked(company=company, score=int(max(0, min(100, round(score)))),
                  product=product or rules["default_product"], reasons=reasons,
                  ask_for=_ask_for(company, obs),
                  decision_maker=bool(obs.get("DECISION_MAKER_TITLE")))


async def rank_companies(state, rules: dict | None = None) -> list[Ranked]:
    """Every company, best first, with its reasons."""
    rules = rules or load_rules()
    async with state._connect() as db:
        db.row_factory = aiosqlite.Row
        async with db.execute(
            """SELECT id, name, domain, website, industry, location, phone, notes,
                      profile_checked_at FROM companies"""
        ) as cursor:
            companies = [dict(r) for r in await cursor.fetchall()]
        # Newest observation per (company, signal, text): people and roles keep
        # one row each; a re-profiled number keeps only its latest value.
        # Observation ids are random; rowid is insertion order, i.e. newest last.
        async with db.execute(
            """SELECT company_id, signal_code, value_num, value_text
               FROM observations o
               WHERE o.rowid IN (SELECT MAX(rowid) FROM observations
                                 GROUP BY company_id, signal_code, COALESCE(value_text, ''))
               ORDER BY o.rowid"""
        ) as cursor:
            rows = [dict(r) for r in await cursor.fetchall()]
    by_company: dict[str, dict[str, list[dict]]] = {}
    for r in rows:
        by_company.setdefault(r["company_id"], {}).setdefault(r["signal_code"], []).append(r)
    ranked = []
    for company in companies:
        company["_profiled"] = bool(company.get("profile_checked_at"))
        ranked.append(score_company(company, by_company.get(company["id"], {}), rules))
    ranked.sort(key=lambda r: (-r.score, r.company.get("industry") or "", r.company.get("name") or ""))
    return ranked


EXPORT_COLUMNS = ("Puesto", "Puntuación", "Empresa", "Ofrecer", "Preguntar por", "Teléfono",
                  "Web", "Sector", "Zona", "Motivos", "Lo que ya sabíamos")


def export_rows(ranked: list[Ranked]) -> list[list]:
    rows = []
    for i, r in enumerate(ranked, start=1):
        c = r.company
        rows.append([i, r.score, c.get("name"), r.product, r.ask_for, c.get("phone") or "",
                     c.get("website") or c.get("domain") or "", c.get("industry"), c.get("location"),
                     "\n".join(r.reasons), c.get("notes") or ""])
    return rows


def write_ranking(ranked: list[Ranked], path: str | Path) -> Path:
    path = Path(path)
    rows = export_rows(ranked)
    if path.suffix.lower() == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Prioridad"
        ws.append(list(EXPORT_COLUMNS))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        widths = {"Empresa": 34, "Ofrecer": 16, "Preguntar por": 34, "Motivos": 60,
                  "Lo que ya sabíamos": 50, "Web": 28, "Sector": 22}
        for i, header in enumerate(EXPORT_COLUMNS, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = widths.get(header, 12)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = ws.dimensions
        wb.save(path)
        return path
    import csv
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, delimiter=";")
        writer.writerow(EXPORT_COLUMNS)
        writer.writerows(rows)
    return path


def summary(ranked: list[Ranked]) -> dict:
    """Counts by product and score band, for the CLI and pages."""
    bands = {"80-100": 0, "60-79": 0, "40-59": 0, "0-39": 0}
    products: dict[str, int] = {}
    for r in ranked:
        key = "80-100" if r.score >= 80 else "60-79" if r.score >= 60 else "40-59" if r.score >= 40 else "0-39"
        bands[key] += 1
        products[r.product or "—"] = products.get(r.product or "—", 0) + 1
    with_person = sum(1 for r in ranked if r.ask_for and "(del CRM)" not in r.ask_for)
    return {"total": len(ranked), "bands": bands, "products": products, "with_person": with_person}


def to_json(ranked: list[Ranked]) -> str:
    return json.dumps([
        {"score": r.score, "product": r.product, "reasons": r.reasons, "ask_for": r.ask_for,
         "name": r.company.get("name"), "industry": r.company.get("industry"),
         "location": r.company.get("location"), "domain": r.company.get("domain")}
        for r in ranked], ensure_ascii=False)
